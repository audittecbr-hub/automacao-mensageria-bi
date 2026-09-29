import openpyxl
import os
import subprocess
import json

exe_path = r'C:\Users\cristhofer.maciel.GRUPOSTUDIO\.vscode\extensions\analysis-services.powerbi-modeling-mcp-0.4.0-win32-x64\server\powerbi-modeling-mcp.exe'

p = subprocess.Popen([exe_path, '--start'], stdin=subprocess.PIPE, stdout=subprocess.PIPE, text=True, encoding='utf-8')
def send(msg):
    p.stdin.write(json.dumps(msg) + '\n')
    p.stdin.flush()
    return json.loads(p.stdout.readline())

send({'jsonrpc': '2.0', 'id': 1, 'method': 'initialize', 'params': {'protocolVersion': '2024-11-05', 'capabilities': {}, 'clientInfo': {'name': 'test', 'version': '1.0'}}})
send({'jsonrpc': '2.0', 'id': 2, 'method': 'tools/call', 'params': {'name': 'connection_operations', 'arguments': {'request': {'operation': 'Connect', 'connectionString': 'Data Source=localhost:57503;Application Name=MCP-Direct'}}}})
dax_res = send({'jsonrpc': '2.0', 'id': 3, 'method': 'tools/call', 'params': {'name': 'dax_query_operations', 'arguments': {'request': {'operation': 'Execute', 'query': 'EVALUATE DISTINCT(vw_powerbi_job_repasse[JOB])'}}}})
p.terminate()

known_jobs = set()
for c in dax_res.get('result', {}).get('content', []):
    if 'resource' in c:
        for line in c['resource']['text'].splitlines()[1:]:
            j = line.strip().strip('"')
            if j:
                known_jobs.add(j)

print(f'Total known jobs in vw_powerbi_job_repasse: {len(known_jobs)}')

desktop = os.path.join(os.environ['USERPROFILE'], 'Desktop')
fpath = os.path.join(desktop, 'Relatorio_Anomalias_Tax_Corporate_Ago_Set_2026_Tratado.xlsx')
wb = openpyxl.load_workbook(fpath, data_only=True)

for sname in ['Painel Repasses', 'Painel Repasse 21.09.26']:
    if sname in wb.sheetnames:
        ws = wb[sname]
        unmapped = []
        for r in range(2, ws.max_row + 1):
            job = str(ws.cell(row=r, column=1).value or '').strip()
            if job and job not in known_jobs:
                unmapped.append((r, job, ws.cell(row=r, column=3).value, ws.cell(row=r, column=7).value))
        print(f'Sheet {sname}: {len(unmapped)} unmapped jobs found.')
        for r, j, cli, val in unmapped:
            print(f'   Row {r}: Job={j}, Cliente={cli}, Valor={val}')
