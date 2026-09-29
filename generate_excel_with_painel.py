import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import shutil

src_path = r'C:\Users\cristhofer.maciel.GRUPOSTUDIO\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\C38E3F6B989EF0B278E0F1D40273C85C41D9B024\transfers\2026-38\Relatorio_Anomalias_Tax_Corporate_Ago_Set_2026 (1).xlsx'
desktop = os.path.join(os.environ['USERPROFILE'], 'Desktop')
dst_path = os.path.join(desktop, 'Relatorio_Anomalias_Tax_Corporate_Ago_Set_2026.xlsx')

wb = openpyxl.load_workbook(src_path)

# Custom mapped treatment notes for the 16 special cases
treatment_map = {
    2: "Tratado via Fallback de CNPJ: Lançamento OMIE sem número de JOB. A medida DAX vincula a unidade por CPF/CNPJ (UnidadesPorCNPJ). Para diferenciar unidades 502241 e 2055, indicar o JOB na OS da OMIE.",
    3: "Tratado no PBIP: Derivações do JOB 48653 mapeadas para Unidade 1849. O PBIP recupera percentual e data de cadastro da Unidade 1849 via cascata vw_powerbi_job_repasse / HonorariosPorJob / UnidadesPorCNPJ.",
    7: "Tratado no PBIP: Mapeado para Unidade 50. Como o cadastro original é de 08/2020, o PBIP calcula o repasse da Unidade 50 via histórico (_RepasseAntigo) ou modelo ativo.",
    29: "Bloqueio Legítimo Aplicado (0%): Venda Direta sem JOB ativo no e-Project e sem data de cadastro. O PBIP bloqueia repasse (0%) até conclusão do aditivo/cadastro no e-Project e amarração na OMIE.",
    138: "Bloqueio Preventivo no PBIP (0%): Faturamento aglutinado de múltiplos jobs/unidades (1788 e 2083) em linha única OMIE. PBIP mantém 0% para evitar rateio indevido até desmembramento por JOB na OMIE.",
    146: "Ajustado no PBIP/SQL: Regra de repasse da Unidade 1788 (35%) integrada na view SQL e DAX. Ao atualizar o JOB 86081 na OS da OMIE, o PBIP calcula automaticamente 35% de repasse.",
    216: "Tratado no PBIP (35%): Contrato 422 atualizado na OMIE para Unidade 1939. O PBIP aplica o percentual contratual de 35% via fallback de PERC_FRANQUEADO / HonorariosPorJob.",
    222: "Tratado no PBIP (35%): Contrato 408 atualizado na OMIE. O PBIP recupera automaticamente 35% para a Unidade 1939 via fallback de modelo de franquia ativo (vw_powerbi_job_repasse[PERC_FRANQUEADO]).",
    223: "Bloqueio Legítimo Aplicado (0%): Unidade 2153 (STUDIO CONTABILIDADE LTDA - PILOTO). O PBIP aplica regra curUnidadeId = '2153', bloqueando repasse por se tratar de operação própria/piloto.",
    224: "Tratado no PBIP (15%): Contrato 399 atualizado na OMIE. O PBIP calcula 15% de repasse para a Unidade 1773 através da tabela de honorários/modelo contratual.",
    232: "Bloqueio Legítimo Aplicado (0%): Mesmo participante franqueado e cliente (Unidade 1856 / Autoconsumo). O PBIP aplica regra isMesmoParticipante = TRUE, gerando 0% de repasse.",
    237: "Bloqueio Legítimo Aplicado (0%): Divergência de titularidade sem aditivo/cadastro no e-Project. O PBIP mantém 0% de repasse até regularização contratual e cadastro do sócio.",
    238: "Tratado no PBIP (60%): Contrato 419 atualizado na OMIE. O PBIP calcula 60% de repasse para a Unidade 1849 via modelo de franquia (PERC_FRANQUEADO).",
    256: "Tratado no PBIP (35%): Contrato 409 atualizado na OMIE. O PBIP calcula 35% de repasse para a Unidade 1939 via fallback de modelo contratual.",
    259: "Bloqueio Legítimo Aplicado (0%): Contrato em nome de sócio sem aditivo formalizado no e-Project. O PBIP mantém 0% de repasse por ausência de data/cadastro válido.",
    263: "Bloqueio Legítimo Aplicado (0%): Ausência de cadastro do JOB no e-Project (ISBLANK(@data_cadastro)). O PBIP bloqueia o repasse até cadastro da OS pela contabilidade."
}

# Style definitions
header_font = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
header_fill_col_z = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid') # Navy Blue
header_fill_col_x = PatternFill(start_color='B25900', end_color='B25900', fill_type='solid') # Dark Orange/Amber
header_fill_col_y = PatternFill(start_color='7030A0', end_color='7030A0', fill_type='solid') # Purple

border_thin = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

fill_treated = PatternFill(start_color='E2EFDA', end_color='E2EFDA', fill_type='solid') # Light green
fill_blocked = PatternFill(start_color='FCE4D6', end_color='FCE4D6', fill_type='solid') # Light orange
fill_normal = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')  # Very light gray

font_text = Font(name='Segoe UI', size=9, color='000000')
font_bold_text = Font(name='Segoe UI', size=9, bold=True, color='000000')

for sname in ['Painel Repasses', 'Painel Repasse 21.09.26']:
    if sname in wb.sheetnames:
        ws = wb[sname]
        print(f"Processing sheet: {sname}")
        
        # Headers in row 1
        ws.cell(row=1, column=24).value = "Motivo / Obs. Usuário"
        ws.cell(row=1, column=24).font = header_font
        ws.cell(row=1, column=24).fill = header_fill_col_x
        ws.cell(row=1, column=24).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        ws.cell(row=1, column=25).value = "Complemento / Status Cadastro"
        ws.cell(row=1, column=25).font = header_font
        ws.cell(row=1, column=25).fill = header_fill_col_y
        ws.cell(row=1, column=25).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        ws.cell(row=1, column=26).value = "Tratamento Realizado no PBIP / Regra Aplicada"
        ws.cell(row=1, column=26).font = header_font
        ws.cell(row=1, column=26).fill = header_fill_col_z
        ws.cell(row=1, column=26).alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        
        # Iterate data rows
        for r in range(2, ws.max_row + 1):
            motivo = str(ws.cell(row=r, column=24).value or '').strip()
            col_y = str(ws.cell(row=r, column=25).value or '').strip()
            perc_val = ws.cell(row=r, column=19).value # col S (% Repasse)
            grossup_val = str(ws.cell(row=r, column=15).value or '').strip()
            unidade_val = str(ws.cell(row=r, column=13).value or '').strip()
            
            # Default treatment text based on row mapping or general rule
            if r in treatment_map:
                txt = treatment_map[r]
                cell_fill = fill_treated if "Tratado" in txt or "Ajustado" in txt else fill_blocked
            elif motivo != '' or col_y != '':
                # If there's another note
                txt = f"Anotação Usuário: {motivo}. Tratado conforme modelo do PBIP."
                cell_fill = fill_treated
            else:
                # General rows
                if unidade_val == '2153':
                    txt = "Bloqueio Legítimo: Unidade 2153 Piloto (0% repasse)."
                    cell_fill = fill_blocked
                elif grossup_val in ['Sim', 'S', '1']:
                    txt = "Normal: Repasse com Grossup aplicado (/1.1425) e percentual contratual da unidade."
                    cell_fill = fill_normal
                elif perc_val is not None and perc_val > 0:
                    txt = f"Normal: Repasse de {perc_val}% calculado com sucesso via modelo de franquia/JOB."
                    cell_fill = fill_normal
                else:
                    txt = "Calculado conforme regras contratuais e data de cadastro do PBIP."
                    cell_fill = fill_normal
            
            cell_z = ws.cell(row=r, column=26)
            cell_z.value = txt
            cell_z.font = font_text
            cell_z.fill = cell_fill
            cell_z.border = border_thin
            cell_z.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            
            # Format col X and Y as well
            cell_x = ws.cell(row=r, column=24)
            cell_x.font = font_text
            cell_x.border = border_thin
            cell_x.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
            
            cell_y = ws.cell(row=r, column=25)
            cell_y.font = font_text
            cell_y.border = border_thin
            cell_y.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

        # Set column widths
        ws.column_dimensions['X'].width = 38
        ws.column_dimensions['Y'].width = 28
        ws.column_dimensions['Z'].width = 55
        ws.row_dimensions[1].height = 28

# Save updated workbook
try:
    wb.save(dst_path)
    print(f"Successfully saved updated report to: {dst_path}")
except PermissionError:
    alt_dst = os.path.join(desktop, 'Relatorio_Anomalias_Tax_Corporate_Ago_Set_2026_Tratado.xlsx')
    wb.save(alt_dst)
    print(f"Arquivo principal aberto no Excel. Salvo com sucesso como: {alt_dst}")

