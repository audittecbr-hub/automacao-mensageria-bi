import openpyxl, os

path = r"C:\Users\cristhofer.maciel.GRUPOSTUDIO\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\C38E3F6B989EF0B278E0F1D40273C85C41D9B024\transfers\2026-38\Relatorio_Anomalias_Tax_Corporate_Ago_Set_2026.xlsx"

print("Tamanho:", os.path.getsize(path), "bytes")
wb = openpyxl.load_workbook(path, data_only=True)
print("Abas:", wb.sheetnames)

for s in wb.sheetnames:
    ws = wb[s]
    print(f"\n--- ABA: {s} (Linhas: {ws.max_row}, Colunas: {ws.max_column}) ---")
    headers = [str(ws.cell(row=4, column=c).value or ws.cell(row=1, column=c).value or '') for c in range(1, min(ws.max_column+1, 30))]
    print("Cabeçalhos:", [h for h in headers if h])
    
    # Check if there is text in columns > 21 (like X, Y, etc.)
    for r in range(1, min(ws.max_row+1, 20)):
        extra_vals = [str(ws.cell(row=r, column=c).value or '') for c in range(21, min(ws.max_column+1, 30))]
        if any(extra_vals):
            print(f"  Linha {r} (Cols 21+): {extra_vals}")
