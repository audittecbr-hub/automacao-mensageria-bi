import openpyxl

path = r'C:\Users\cristhofer.maciel.GRUPOSTUDIO\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\C38E3F6B989EF0B278E0F1D40273C85C41D9B024\transfers\2026-38\Relatorio_Anomalias_Tax_Corporate_Ago_Set_2026 (1).xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
ws = wb['Painel Repasses']

headers = [ws.cell(1, c).value for c in range(1, ws.max_column+1)]
print('Headers:', headers)

rows_with_motivo = []
for r in range(2, ws.max_row+1):
    motivo = ws.cell(r, 24).value
    col_y = ws.cell(r, 25).value
    if motivo or col_y:
        row_data = {
            'row': r,
            'job': ws.cell(r, 1).value,
            'data_cad': str(ws.cell(r, 2).value),
            'cliente': ws.cell(r, 3).value,
            'bandeira': ws.cell(r, 4).value,
            'cnpj': ws.cell(r, 5).value,
            'categoria': ws.cell(r, 6).value,
            'valor': ws.cell(r, 7).value,
            'nf': ws.cell(r, 8).value,
            'unidade': ws.cell(r, 12).value,
            'nome_unidade': ws.cell(r, 13).value,
            'grossup': ws.cell(r, 14).value,
            'perc': ws.cell(r, 18).value,
            'motivo': motivo,
            'col_y': col_y
        }
        rows_with_motivo.append(row_data)

print(f'Total rows with motivo/col_y: {len(rows_with_motivo)}')
for item in rows_with_motivo:
    print(f"Row {item['row']}: Job={item['job']}, Unid={item['unidade']}, CNPJ={item['cnpj']}, Cliente={item['cliente']} | Perc={item['perc']} | Motivo={item['motivo']} | Col_Y={item['col_y']}")

wb.close()
