import pandas as pd
import numpy as np
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

df = pd.read_csv("raw_metas_aug_sep.csv", encoding="utf-8")

def clean_str(val):
    if pd.isna(val):
        return ""
    return str(val).strip()

def clean_num(val):
    if pd.isna(val):
        return 0.0
    s = str(val).replace(".", "").replace(",", ".")
    try:
        return float(s)
    except:
        return 0.0

def clean_date_str(val):
    if pd.isna(val) or not str(val).strip():
        return ""
    s = str(val).strip().split(" ")[0]
    return s

records = []
for idx, r in df.iterrows():
    cod_omie = clean_str(r['Metas[codigo_lancamento_omie]'])
    dt_emissao = clean_date_str(r['Metas[data_emissao]'])
    dt_lancamento = clean_date_str(r['Metas[data_lancamento]'])
    dt_vencimento = clean_date_str(r['Metas[data_vencimento]'])
    categoria = clean_str(r['Metas[categoria]']).upper()
    desc_cat = clean_str(r['Metas[descricao_cat]'])
    desc_dept = clean_str(r['Metas[descricao_dept]'])
    bandeira = clean_str(r['Metas[bandeira]'])
    cnpj_cpf = clean_str(r['Metas[cnpj_cpf]'])
    razao_social = clean_str(r['Metas[razao_social]'])
    valor_conta = clean_num(r['Metas[valor_conta]'])
    
    job = clean_str(r['Metas[numero_contrato]'])
    if not job or job in ["nan", "NaN", "None", "0"]:
        job_lookup = clean_str(r['[Job_Contrato]'])
        if job_lookup and job_lookup not in ["nan", "NaN", "None", "0"]:
            job = job_lookup
        else:
            job = ""
            
    perc_job = clean_num(r['[Vw_Perc_Job]'])
    perc_tbl = clean_num(r['[Tbl_Honorario]'])
    honorario = perc_job if perc_job > 0 else perc_tbl
    
    unidade_id = clean_str(r['[Vw_Unidade_Id]'])
    if unidade_id.endswith(".0"):
        unidade_id = unidade_id[:-2]
    unidade_nome = clean_str(r['[Vw_Unidade_Nome]'])
    rede = clean_str(r['[Vw_Rede_Distribuicao]'])
    rede_old = clean_str(r['[Vw_Rede_Distribuicao_Old]'])
    franq = clean_str(r['[Vw_Part_Franqueado]'])
    cli = clean_str(r['[Vw_Part_Cliente]'])
    data_cad = clean_date_str(r['[Vw_Data_Cadastro]'])
    grossup = clean_str(r['[Vw_Grossup]'])
    retencao = clean_num(r['[Vw_Retencao]'])
    
    # Check legitimate rules
    rede_upper = rede.upper()
    rede_old_upper = rede_old.upper()
    unidade_nome_upper = unidade_nome.upper()
    franq_upper = franq.upper().strip()
    cli_upper = cli.upper().strip()
    
    is_store_xp = ("STORE" in rede_upper) or ("XP" in rede_upper) or ("STORE" in rede_old_upper) or ("XP" in rede_old_upper)
    is_unidade_2153 = (unidade_id == "2153") or ("STUDIO CONTABILIDADE LTDA - PILOTO" in unidade_nome_upper)
    is_mesmo_part = bool(franq_upper and cli_upper and (franq_upper == cli_upper))
    
    # Flags
    has_job = bool(job and job not in ["0", "-", "NAN", "NONE"])
    job_found_in_view = bool(clean_str(r['[Vw_Unidade_Id]']) or perc_job > 0 or clean_str(r['[Vw_Data_Cadastro]']))
    
    # Classification
    if not has_job:
        status = "1. Sem JOB / Contrato"
        motivo = "Campo numero_contrato em branco ou nulo no lançamento Omie."
        tipo_grupo = "Anomalia: Sem JOB"
    elif is_store_xp:
        status = "Regra Legítima: Bloqueio Rede Store/XP"
        motivo = f"Rede identificada como Store/XP ({rede or rede_old}). Repasse zerado conforme regra de negócio."
        tipo_grupo = "Regra Legítima (Sem Repasse)"
    elif is_unidade_2153:
        status = "Regra Legítima: Bloqueio Unidade 2153 Piloto"
        motivo = "Unidade 2153 (STUDIO CONTABILIDADE LTDA - PILOTO). Repasse zerado conforme regra de negócio."
        tipo_grupo = "Regra Legítima (Sem Repasse)"
    elif is_mesmo_part:
        status = "Regra Legítima: Bloqueio Franqueado = Cliente"
        motivo = f"Autoconsumo: Franqueado ({franq}) é o mesmo que o Cliente ({cli})."
        tipo_grupo = "Regra Legítima (Sem Repasse)"
    elif not job_found_in_view and perc_tbl == 0:
        status = "2. JOB não cadastrado na View"
        motivo = f"Número do contrato '{job}' não localizado em vw_powerbi_job_repasse nem na tabela de honorários."
        tipo_grupo = "Anomalia: JOB Inexistente no Cadastro"
    elif honorario == 0:
        status = "3. Percentual Zerado sem Regra"
        motivo = "JOB localizado na base, porém o percentual de repasse está cadastrado como 0,00% sem enquadrar em regra de bloqueio."
        tipo_grupo = "Anomalia: Percentual 0%"
    elif not data_cad:
        status = "4. JOB sem Data de Cadastro"
        motivo = "JOB possui percentual de repasse, mas a Data de Cadastro está vazia na View (o que anula o repasse no DAX)."
        tipo_grupo = "Anomalia: Data Cadastro Vazia"
    elif not unidade_id:
        status = "5. Sem Unidade Identificada"
        motivo = "JOB localizado, mas a Unidade ID não foi identificada."
        tipo_grupo = "Anomalia: Sem Unidade"
    else:
        status = "Normal / Repasse Elegível"
        motivo = "Lançamento válido com JOB, Unidade e Percentual calculados normalmente."
        tipo_grupo = "Normal"
        
    records.append({
        "Cod_Omie": cod_omie,
        "Data_Emissao": dt_emissao,
        "Data_Lancamento": dt_lancamento,
        "Data_Vencimento": dt_vencimento,
        "Categoria": categoria,
        "Descricao_Categoria": desc_cat,
        "Descricao_Departamento": desc_dept,
        "Bandeira": bandeira,
        "CNPJ_CPF": cnpj_cpf,
        "Razao_Social": razao_social,
        "Valor_Conta_R$": valor_conta,
        "Numero_Contrato_JOB": job,
        "Perc_Honorario_JOB_%": perc_job,
        "Perc_Honorario_Tabela_%": perc_tbl,
        "Perc_Honorario_Final_%": honorario,
        "Unidade_ID": unidade_id,
        "Unidade_Nome": unidade_nome,
        "Rede_Distribuicao": rede,
        "Rede_Distribuicao_Old": rede_old,
        "Participante_Franqueado": franq,
        "Participante_Cliente": cli,
        "Data_Cadastro_JOB": data_cad,
        "Grossup": grossup,
        "Retencao_%": retencao,
        "Classificacao": status,
        "Grupo_Status": tipo_grupo,
        "Motivo_Detalhamento": motivo
    })

df_final = pd.DataFrame(records)

# Filter subsets
df_sem_job = df_final[df_final["Classificacao"] == "1. Sem JOB / Contrato"].copy()
df_perc_zero = df_final[df_final["Classificacao"] == "3. Percentual Zerado sem Regra"].copy()
df_job_nao_cad = df_final[df_final["Classificacao"] == "2. JOB não cadastrado na View"].copy()
df_sem_data_cad = df_final[df_final["Classificacao"] == "4. JOB sem Data de Cadastro"].copy()
df_regras_legitimas = df_final[df_final["Grupo_Status"] == "Regra Legítima (Sem Repasse)"].copy()
df_todas_anomalias = df_final[df_final["Grupo_Status"].str.startswith("Anomalia") | (df_final["Classificacao"] == "1. Sem JOB / Contrato")].copy()

output_excel = "Relatorio_Anomalias_Repasse_Ago_Set_2026.xlsx"
wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Styles
font_title = Font(name="Segoe UI", size=14, bold=True, color="FFFFFF")
font_subtitle = Font(name="Segoe UI", size=10, italic=True, color="D4AF37")
font_section = Font(name="Segoe UI", size=11, bold=True, color="0D1120")
font_header = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
font_bold = Font(name="Segoe UI", size=10, bold=True, color="0D1120")
font_regular = Font(name="Segoe UI", size=9, color="1F2937")
font_small = Font(name="Segoe UI", size=8, color="6B7280")

fill_title = PatternFill(start_color="0D1120", end_color="0D1120", fill_type="solid")
fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
fill_header_gold = PatternFill(start_color="854D0E", end_color="854D0E", fill_type="solid")
fill_header_red = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
fill_header_orange = PatternFill(start_color="9A3412", end_color="9A3412", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_card_kpi = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
fill_alert_red = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

border_thin = Side(style="thin", color="CBD5E1")
border_card = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)
border_bottom_thick = Border(bottom=Side(style="medium", color="D4AF37"))

align_left = Alignment(horizontal="left", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ----------------------------------------------------
# 1. TAB: RESUMO EXECUTIVO
# ----------------------------------------------------
ws_resumo = wb.create_sheet(title="Resumo Executivo")
ws_resumo.views.sheetView[0].showGridLines = True

# Title Header
ws_resumo.merge_cells("A1:G1")
ws_resumo["A1"] = "RELATÓRIO EXECUTIVO DE AUDITORIA E ANOMALIAS DE REPASSE"
ws_resumo["A1"].font = font_title
ws_resumo["A1"].fill = fill_title
ws_resumo["A1"].alignment = align_center

ws_resumo.merge_cells("A2:G2")
ws_resumo["A2"] = "Corte Temporal: Agosto e Setembro/2026 | Análise de Contratos, Percentuais e Regras de Negócio"
ws_resumo["A2"].font = font_subtitle
ws_resumo["A2"].fill = fill_title
ws_resumo["A2"].alignment = align_center
ws_resumo.row_dimensions[1].height = 28
ws_resumo.row_dimensions[2].height = 18

# Summary KPI Cards
kpis = [
    ("Total de Lançamentos", len(df_final), f"R$ {df_final['Valor_Conta_R$'].sum():,.2f}"),
    ("Lançamentos com Anomalias", len(df_todas_anomalias), f"R$ {df_todas_anomalias['Valor_Conta_R$'].sum():,.2f}"),
    ("1. Sem JOB / Contrato", len(df_sem_job), f"R$ {df_sem_job['Valor_Conta_R$'].sum():,.2f}"),
    ("2. JOB Não Cadastrado", len(df_job_nao_cad), f"R$ {df_job_nao_cad['Valor_Conta_R$'].sum():,.2f}"),
    ("3. Percentual Zerado (Sem Regra)", len(df_perc_zero), f"R$ {df_perc_zero['Valor_Conta_R$'].sum():,.2f}"),
    ("4. JOB sem Data Cadastro", len(df_sem_data_cad), f"R$ {df_sem_data_cad['Valor_Conta_R$'].sum():,.2f}"),
    ("Regras Legítimas de Bloqueio", len(df_regras_legitimas), f"R$ {df_regras_legitimas['Valor_Conta_R$'].sum():,.2f}"),
    ("Lançamentos Normais / OK", len(df_final[df_final['Grupo_Status'] == 'Normal']), f"R$ {df_final[df_final['Grupo_Status'] == 'Normal']['Valor_Conta_R$'].sum():,.2f}")
]

row_kpi = 4
ws_resumo.cell(row=row_kpi, column=1, value="PAINEL CONSOLIDADO DE INDICADORES").font = font_section
row_kpi += 1

ws_resumo.cell(row=row_kpi, column=1, value="Indicador / Classificação").font = font_header
ws_resumo.cell(row=row_kpi, column=1).fill = fill_header
ws_resumo.cell(row=row_kpi, column=1).alignment = align_left

ws_resumo.cell(row=row_kpi, column=2, value="Qtd Lançamentos").font = font_header
ws_resumo.cell(row=row_kpi, column=2).fill = fill_header
ws_resumo.cell(row=row_kpi, column=2).alignment = align_right

ws_resumo.cell(row=row_kpi, column=3, value="% do Total").font = font_header
ws_resumo.cell(row=row_kpi, column=3).fill = fill_header
ws_resumo.cell(row=row_kpi, column=3).alignment = align_right

ws_resumo.cell(row=row_kpi, column=4, value="Valor Total Faturado (R$)").font = font_header
ws_resumo.cell(row=row_kpi, column=4).fill = fill_header
ws_resumo.cell(row=row_kpi, column=4).alignment = align_right

ws_resumo.cell(row=row_kpi, column=5, value="Status / Ação Recomendada").font = font_header
ws_resumo.cell(row=row_kpi, column=5).fill = fill_header
ws_resumo.cell(row=row_kpi, column=5).alignment = align_left

row_curr = row_kpi + 1
total_qtd = len(df_final)

for label, qtd, val_str in kpis:
    pct = (qtd / total_qtd) if total_qtd > 0 else 0
    val_num = df_final[df_final['Classificacao'] == label]['Valor_Conta_R$'].sum() if label in df_final['Classificacao'].values else (
        df_todas_anomalias['Valor_Conta_R$'].sum() if "Anomalias" in label else (
            df_regras_legitimas['Valor_Conta_R$'].sum() if "Regras Legítimas" in label else (
                df_final[df_final['Grupo_Status'] == 'Normal']['Valor_Conta_R$'].sum() if "Normais" in label else df_final['Valor_Conta_R$'].sum()
            )
        )
    )
    
    acao = "Auditar e preencher contrato no Omie" if "Sem JOB" in label else (
        "Cadastrar contrato no ERP/View de repasse" if "Não Cadastrado" in label else (
            "Verificar alíquota/modelo de honorário no contrato" if "Percentual Zerado" in label else (
                "Preencher data de cadastro da vigência no contrato" if "sem Data" in label else (
                    "Bloqueio previsto pelas 3 regras ativas (Store/XP, 2153, Autoconsumo)" if "Legítimas" in label else (
                        "Atenção prioritária para correção operacional" if "Anomalias" in label else "Base regular e operacional"
                    )
                )
            )
        )
    )
    
    ws_resumo.cell(row=row_curr, column=1, value=label).font = font_bold if "Total" in label or "Anomalias" in label else font_regular
    ws_resumo.cell(row=row_curr, column=1).alignment = align_left
    ws_resumo.cell(row=row_curr, column=1).border = border_card
    
    c2 = ws_resumo.cell(row=row_curr, column=2, value=qtd)
    c2.font = font_bold if "Total" in label or "Anomalias" in label else font_regular
    c2.alignment = align_right
    c2.number_format = "#,##0"
    c2.border = border_card
    
    c3 = ws_resumo.cell(row=row_curr, column=3, value=pct)
    c3.font = font_bold if "Total" in label or "Anomalias" in label else font_regular
    c3.alignment = align_right
    c3.number_format = "0.0%"
    c3.border = border_card
    
    c4 = ws_resumo.cell(row=row_curr, column=4, value=val_num)
    c4.font = font_bold if "Total" in label or "Anomalias" in label else font_regular
    c4.alignment = align_right
    c4.number_format = "R$ #,##0.00"
    c4.border = border_card
    
    c5 = ws_resumo.cell(row=row_curr, column=5, value=acao)
    c5.font = font_regular
    c5.alignment = align_left
    c5.border = border_card
    
    if "Anomalias" in label:
        for c in range(1, 6):
            ws_resumo.cell(row=row_curr, column=c).fill = fill_alert_red
            
    row_curr += 1

# Category Breakdown Table
row_curr += 2
ws_resumo.cell(row=row_curr, column=1, value="DISTRIBUIÇÃO DE ANOMALIAS POR CATEGORIA").font = font_section
row_curr += 1

headers_cat = ["Categoria", "Total Lançamentos", "Sem JOB / Contrato", "JOB Não Cadastrado", "Percentual Zerado", "JOB sem Data Cad.", "Valor Total (R$)"]
for col_i, h in enumerate(headers_cat, 1):
    c = ws_resumo.cell(row=row_curr, column=col_i, value=h)
    c.font = font_header
    c.fill = fill_header
    c.alignment = align_header

row_curr += 1
for cat, grp in df_final.groupby("Categoria"):
    total_c = len(grp)
    s_job = len(grp[grp["Classificacao"] == "1. Sem JOB / Contrato"])
    n_cad = len(grp[grp["Classificacao"] == "2. JOB não cadastrado na View"])
    p_zero = len(grp[grp["Classificacao"] == "3. Percentual Zerado sem Regra"])
    s_data = len(grp[grp["Classificacao"] == "4. JOB sem Data de Cadastro"])
    val_c = grp["Valor_Conta_R$"].sum()
    
    ws_resumo.cell(row=row_curr, column=1, value=cat).font = font_bold
    ws_resumo.cell(row=row_curr, column=1).alignment = align_left
    ws_resumo.cell(row=row_curr, column=1).border = border_card
    
    for c_i, val in enumerate([total_c, s_job, n_cad, p_zero, s_data], 2):
        c = ws_resumo.cell(row=row_curr, column=c_i, value=val)
        c.font = font_regular
        c.alignment = align_right
        c.number_format = "#,##0"
        c.border = border_card
        
    c_val = ws_resumo.cell(row=row_curr, column=7, value=val_c)
    c_val.font = font_regular
    c_val.alignment = align_right
    c_val.number_format = "R$ #,##0.00"
    c_val.border = border_card
    
    row_curr += 1

# Format column widths for Resumo
ws_resumo.column_dimensions["A"].width = 38
ws_resumo.column_dimensions["B"].width = 18
ws_resumo.column_dimensions["C"].width = 14
ws_resumo.column_dimensions["D"].width = 24
ws_resumo.column_dimensions["E"].width = 45
ws_resumo.column_dimensions["F"].width = 20
ws_resumo.column_dimensions["G"].width = 22

# ----------------------------------------------------
# Helper to write data tables in tabs
# ----------------------------------------------------
def populate_data_sheet(ws, dataframe, title_text, header_fill_color):
    ws.views.sheetView[0].showGridLines = True
    
    # Title
    ws.merge_cells("A1:N1")
    ws["A1"] = title_text
    ws["A1"].font = font_title
    ws["A1"].fill = fill_title
    ws["A1"].alignment = align_center
    
    ws.merge_cells("A2:N2")
    ws["A2"] = f"Total de Registros: {len(dataframe)} | Valor Consolidado: R$ {dataframe['Valor_Conta_R$'].sum():,.2f}"
    ws["A2"].font = font_subtitle
    ws["A2"].fill = fill_title
    ws["A2"].alignment = align_center
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 18
    
    columns_to_show = [
        ("Cód. Omie", "Cod_Omie", "center", None),
        ("Dt Emissão", "Data_Emissao", "center", None),
        ("Dt Lançamento", "Data_Lancamento", "center", None),
        ("Categoria", "Categoria", "left", None),
        ("Bandeira", "Bandeira", "left", None),
        ("CNPJ / CPF", "CNPJ_CPF", "center", None),
        ("Razão Social Cliente", "Razao_Social", "left", None),
        ("Valor Conta (R$)", "Valor_Conta_R$", "right", "R$ #,##0.00"),
        ("Nº Contrato / JOB", "Numero_Contrato_JOB", "center", None),
        ("% Repasse Final", "Perc_Honorario_Final_%", "right", "0.00%"),
        ("ID Unidade", "Unidade_ID", "center", None),
        ("Nome Unidade", "Unidade_Nome", "left", None),
        ("Rede Distribuição", "Rede_Distribuicao", "left", None),
        ("Rede Antiga", "Rede_Distribuicao_Old", "left", None),
        ("Part. Franqueado", "Participante_Franqueado", "left", None),
        ("Part. Cliente", "Participante_Cliente", "left", None),
        ("Dt Cadastro JOB", "Data_Cadastro_JOB", "center", None),
        ("Grossup", "Grossup", "center", None),
        ("Retenção %", "Retencao_%", "right", "0.00%"),
        ("Classificação", "Classificacao", "left", None),
        ("Detalhamento do Motivo", "Motivo_Detalhamento", "left", None),
    ]
    
    # Headers
    header_row = 4
    ws.row_dimensions[header_row].height = 24
    for col_idx, (header_label, col_key, align_type, num_fmt) in enumerate(columns_to_show, 1):
        c = ws.cell(row=header_row, column=col_idx, value=header_label)
        c.font = font_header
        c.fill = header_fill_color
        c.alignment = align_header
        c.border = border_card
        
    # Data rows
    cur_row = header_row + 1
    for r_idx, (_, row_data) in enumerate(dataframe.iterrows()):
        ws.row_dimensions[cur_row].height = 19
        zebra = (r_idx % 2 == 1)
        
        for col_idx, (header_label, col_key, align_type, num_fmt) in enumerate(columns_to_show, 1):
            val = row_data[col_key]
            c = ws.cell(row=cur_row, column=col_idx, value=val)
            c.font = font_regular
            c.border = border_card
            
            if zebra:
                c.fill = fill_zebra
                
            if align_type == "right":
                c.alignment = align_right
            elif align_type == "center":
                c.alignment = align_center
            else:
                c.alignment = align_left
                
            if num_fmt:
                c.number_format = num_fmt
                if "%" in num_fmt and isinstance(val, (int, float)):
                    c.value = val / 100.0 if val > 1.0 else val
                    
        cur_row += 1
                    
    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in [1, 2]:
                continue
            v = str(cell.value or "")
            if len(v) > max_len:
                max_len = len(v)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    # Freeze header
    ws.freeze_panes = ws.cell(row=5, column=1)

# Tab 2: Sem JOB ou Contrato
ws_sem_job = wb.create_sheet(title="Sem JOB ou Contrato")
populate_data_sheet(ws_sem_job, df_sem_job, "LANÇAMENTOS SEM JOB / CONTRATO (CAMPO EM BRANCO)", fill_header_red)

# Tab 3: Percentual Zerado sem Regra
ws_perc_zero = wb.create_sheet(title="Percentual 0% sem Regra")
populate_data_sheet(ws_perc_zero, df_perc_zero, "LANÇAMENTOS COM PERCENTUAL ZERADO (SEM REGRA DE BLOQUEIO)", fill_header_orange)

# Tab 4: JOB Não Cadastrado na View
ws_job_nao_cad = wb.create_sheet(title="JOB Nao Cadastrado")
populate_data_sheet(ws_job_nao_cad, df_job_nao_cad, "LANÇAMENTOS COM CONTRATO NÃO LOCALIZADO NA BASE DE JOBS", fill_header_gold)

# Tab 5: JOB sem Data Cadastro
ws_sem_data_cad = wb.create_sheet(title="JOB sem Data Cad")
populate_data_sheet(ws_sem_data_cad, df_sem_data_cad, "LANÇAMENTOS COM PERCENTUAL MAS SEM DATA DE CADASTRO", fill_header_orange)

# Tab 6: Regras Legítimas de Bloqueio
ws_regras = wb.create_sheet(title="Regras Legitimas Bloqueio")
populate_data_sheet(ws_regras, df_regras_legitimas, "LANÇAMENTOS ISENTOS DE REPASSE POR REGRA DE NEGÓCIO", fill_header)

# Tab 7: Todas as Anomalias
ws_anomalias = wb.create_sheet(title="Todas Inconsistencias")
populate_data_sheet(ws_anomalias, df_todas_anomalias, "CONSOLIDADO DE TODAS AS ANOMALIAS E INCONSISTÊNCIAS IDENTIFICADAS", fill_header_red)

# Tab 8: Base Completa
ws_completa = wb.create_sheet(title="Base Completa")
populate_data_sheet(ws_completa, df_final, "BASE COMPLETA DE LANÇAMENTOS (AGOSTO E SETEMBRO/2026)", fill_header)

wb.save(output_excel)
print(f"Excel report successfully generated: {output_excel}")
